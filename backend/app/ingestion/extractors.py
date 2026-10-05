"""Text extraction per format.

Each extractor returns `ExtractedDocument`: the text plus, where the format genuinely has
pages, a page map so chunks can carry real page numbers. Markdown and plain text have no
pages and report none — the system never fabricates a page number (README §8.4).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from app.core.errors import ExtractionError


@dataclass
class ExtractedDocument:
    text: str
    page_count: int | None = None
    # (char_offset_start, page_number) markers, ascending — used to map a chunk to pages.
    page_offsets: list[tuple[int, int]] = field(default_factory=list)

    def page_for_offset(self, offset: int) -> int | None:
        if not self.page_offsets:
            return None
        page = self.page_offsets[0][1]
        for start, number in self.page_offsets:
            if start <= offset:
                page = number
            else:
                break
        return page


def extract_markdown(data: bytes) -> ExtractedDocument:
    return ExtractedDocument(text=_decode(data))


def extract_text(data: bytes) -> ExtractedDocument:
    return ExtractedDocument(text=_decode(data))


def extract_pdf(data: bytes, max_pages: int) -> ExtractedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # pypdf raises a variety of parse errors
        raise ExtractionError(f"The PDF could not be opened: {exc}") from exc

    if len(reader.pages) > max_pages:
        raise ExtractionError(
            f"The PDF has {len(reader.pages)} pages; the limit is {max_pages}.",
            page_count=len(reader.pages),
        )

    parts: list[str] = []
    offsets: list[tuple[int, int]] = []
    cursor = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        page_text = page_text.strip()
        if not page_text:
            continue
        offsets.append((cursor, number))
        block = page_text + "\n\n"
        parts.append(block)
        cursor += len(block)

    text = "".join(parts)
    if not text.strip():
        raise ExtractionError(
            "No selectable text was found. Scanned PDFs need OCR, which this service does not perform."
        )
    return ExtractedDocument(text=text, page_count=len(reader.pages), page_offsets=offsets)


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"The DOCX file could not be opened: {exc}") from exc

    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = (para.style.name or "").lower() if para.style else ""
        # Map Word heading styles onto markdown so the structure-aware chunker sees them.
        match = re.match(r"heading (\d)", style)
        if match:
            level = min(int(match.group(1)), 4)
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                lines.append("| " + " | ".join(cells) + " |")
        lines.append("")

    text = "\n".join(lines)
    if not text.strip():
        raise ExtractionError("The DOCX file contains no readable text.")
    return ExtractedDocument(text=text)


def extract(data: bytes, extension: str, max_pdf_pages: int = 500) -> ExtractedDocument:
    if extension == ".pdf":
        return extract_pdf(data, max_pdf_pages)
    if extension == ".docx":
        return extract_docx(data)
    if extension == ".md":
        return extract_markdown(data)
    if extension == ".txt":
        return extract_text(data)
    raise ExtractionError(f"No extractor for {extension}.")


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")
